import openpyxl


def get_flight_search_data():
    workbook = openpyxl.load_workbook(
        "testdata/FlightSearchData.xlsx"
    )

    sheet = workbook["FlightSearchData"]

    data = []

    for row in range(2, sheet.max_row + 1):
        from_city = sheet.cell(row=row, column=1).value
        to_city = sheet.cell(row=row, column=2).value
        departure_date = sheet.cell(row=row, column=3).value

        if from_city and to_city and departure_date:
            data.append(
                (from_city, to_city, departure_date)
            )

    workbook.close()

    return data